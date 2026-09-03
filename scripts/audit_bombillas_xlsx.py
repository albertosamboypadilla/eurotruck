from pathlib import Path
from collections import Counter
import json
from openpyxl import load_workbook

src = Path('/home/ubuntu/upload/Bombillas_DieselTechnic.xlsx')
out = Path('/home/ubuntu/eurotruck-clone/validation/bombillas-xlsx-audit.json')
wb = load_workbook(src, read_only=True, data_only=True)
report = {'file': src.name, 'sheets': []}
for ws in wb.worksheets:
    rows = ws.iter_rows(values_only=True)
    first = next(rows, ())
    headers = [str(v).strip() if v is not None else '' for v in first]
    records = []
    for row in rows:
        vals = list(row)
        if any(v not in (None, '') for v in vals):
            records.append(vals)
    report['sheets'].append({
        'title': ws.title,
        'max_row': ws.max_row,
        'max_column': ws.max_column,
        'headers': headers,
        'non_empty_rows': len(records),
        'sample_rows': records[:5],
    })
report['defined_names'] = list(wb.defined_names.keys())
out.parent.mkdir(parents=True, exist_ok=True)
out.write_text(json.dumps(report, ensure_ascii=False, indent=2, default=str), encoding='utf-8')
print(json.dumps({'sheets': len(report['sheets']), 'report': str(out), 'details': report['sheets']}, ensure_ascii=False, default=str))
