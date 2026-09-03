import json
from pathlib import Path
from openpyxl import load_workbook

catalog = json.loads(Path('/tmp/eurotruck-assets/catalog.json').read_text(encoding='utf-8'))
print('catalog_type', type(catalog).__name__, 'count', len(catalog))
print('catalog_keys', sorted(catalog[0].keys()))
print('catalog_sample', json.dumps(catalog[0], ensure_ascii=False, indent=2)[:2400])
wb = load_workbook('/home/ubuntu/upload/Bombillas_DieselTechnic.xlsx', read_only=True, data_only=True)
ws = wb.active
print('xlsx_headers', [c.value for c in next(ws.iter_rows())])
