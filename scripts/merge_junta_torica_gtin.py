import json
import re
from pathlib import Path
from openpyxl import load_workbook

excel = Path('/home/ubuntu/upload/Junta_Torica_DieselTechnic_1.xlsx')
catalog_path = Path('/home/ubuntu/webdev-static-assets/catalog-with-piston-rings-20260904.json')
map_path = Path('/home/ubuntu/webdev-static-assets/diesel-gtin-map-piston-rings-20260904.json')
out_catalog = Path('/home/ubuntu/webdev-static-assets/catalog-with-junta-torica-gtin-20260904.json')
out_map = Path('/home/ubuntu/webdev-static-assets/diesel-gtin-map-junta-torica-20260904.json')

rows = list(load_workbook(excel, read_only=True, data_only=True).active.iter_rows(min_row=2, values_only=True))
source = {}
for description, brand, part, gtin_cell in rows:
    sku = str(part or '').strip()
    gtins = sorted({value.strip() for value in str(gtin_cell or '').replace(',', '|').split('|') if value.strip().isdigit()})
    if sku and gtins:
        source[sku] = {'description': str(description or '').strip(), 'brand': str(brand or '').strip(), 'gtins': gtins}

items = json.loads(catalog_path.read_text())
gtin_map = json.loads(map_path.read_text())
by_sku = {str(item.get('sku') or '').strip(): item for item in items}

def next_internal_code():
    values = [int(str(item.get('internalCode') or '')) for item in items if str(item.get('internalCode') or '').isdigit()]
    return max(values or [0]) + 1

def slugify(value):
    value = re.sub(r'[^a-z0-9]+', '-', value.lower()).strip('-')
    return value or 'junta-torica'

updated = 0
added = []
for sku, data in source.items():
    current = by_sku.get(sku)
    existing = set((current or {}).get('gtins') or []) | set((gtin_map.get(sku) or {}).get('gtins') or [])
    merged = sorted(existing | set(data['gtins']))
    if current:
        current['gtins'] = merged
        current['barcode'] = current.get('barcode') or merged[0]
        updated += 1
    else:
        code = f'{next_internal_code():05d}'
        slug = f"dt-{sku.replace('.', '-')}-{slugify(data['description'])}"
        current = {
            'id': f'junta-torica-{sku}', 'sku': sku, 'name': data['description'],
            'description': f"DT Spare Parts {sku} {data['description']}", 'slug': slug,
            'url': f'https://partnerportal.dieseltechnic.com/es/products/{slug}',
            'image': '', 'imageFull': '', 'imageSource': '', 'application': '', 'applicationId': '',
            'brand': data['brand'] or 'DT Spare Parts', 'brandId': 'DT', 'manufacturer': '',
            'category': 'Motor', 'subcategory': 'Junta tórica', 'usage': [], 'replaces': '', 'mainOe': '',
            'packagingAmount': 1, 'salesUnit': 'PCS', 'badges': [], 'isProductNews': False,
            'isProductPromotion': False, 'crossReferences': [], 'internalCode': code,
            'barcode': merged[0], 'gtins': merged,
        }
        items.append(current)
        by_sku[sku] = current
        added.append(sku)
    gtin_map[sku] = {'sku': sku, 'gtins': merged}

out_catalog.write_text(json.dumps(items, ensure_ascii=False, separators=(',', ':')))
out_map.write_text(json.dumps(gtin_map, ensure_ascii=False, separators=(',', ':')))
print(json.dumps({'source_rows': len(rows), 'unique_source': len(source), 'updated_existing': updated, 'added_missing': len(added), 'added_skus': added, 'catalog_count': len(items), 'map_count': len(gtin_map), 'catalog_path': str(out_catalog), 'map_path': str(out_map)}, ensure_ascii=False))
