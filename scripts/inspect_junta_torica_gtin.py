import json
from pathlib import Path
from openpyxl import load_workbook

path = Path('/home/ubuntu/upload/Junta_Torica_DieselTechnic_1.xlsx')
workbook = load_workbook(path, read_only=True, data_only=True)
result = {}
for sheet in workbook.worksheets:
    rows = list(sheet.iter_rows(values_only=True))
    result[sheet.title] = {
        'dimensions': sheet.max_row,
        'columns': list(rows[0]) if rows else [],
        'sample': [list(row) for row in rows[1:6]],
        'non_empty_rows': [list(row) for row in rows[1:] if any(value not in (None, '') for value in row)],
    }
print(json.dumps({name: {**data, 'non_empty_count': len(data['non_empty_rows'])} for name, data in result.items()}, ensure_ascii=False, indent=2))
