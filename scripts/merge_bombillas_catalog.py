from pathlib import Path
from openpyxl import load_workbook
import json, re, unicodedata

CATALOG_IN = Path('/tmp/eurotruck-assets/catalog.json')
GTIN_IN = Path('/tmp/eurotruck-assets/gtin.json')
XLSX = Path('/home/ubuntu/upload/Bombillas_DieselTechnic.xlsx')
ASSET_DIR = Path('/home/ubuntu/webdev-static-assets')
ASSET_DIR.mkdir(parents=True, exist_ok=True)
CATALOG_OUT = ASSET_DIR / 'diesel-catalog-bombillas-merged.json'
GTIN_OUT = ASSET_DIR / 'diesel-gtin-map-bombillas-merged.json'
REPORT_OUT = Path('/home/ubuntu/eurotruck-clone/validation/bombillas-import-report.json')


def norm_sku(value):
    if value is None:
        return ''
    s = str(value).strip().replace(',', '.')
    s = re.sub(r'^(?:DT\s*[- ]?)', '', s, flags=re.I)
    return s


def norm_gtin(value):
    return re.sub(r'\D', '', str(value or ''))


def slugify(value):
    raw = unicodedata.normalize('NFKD', value).encode('ascii', 'ignore').decode('ascii').lower()
    return re.sub(r'[^a-z0-9]+', '-', raw).strip('-')

catalog = json.loads(CATALOG_IN.read_text(encoding='utf-8'))
gtin_map = json.loads(GTIN_IN.read_text(encoding='utf-8'))
by_sku = {norm_sku(p.get('sku')): p for p in catalog}
wb = load_workbook(XLSX, read_only=True, data_only=True)
ws = wb.active
headers = [str(c.value).strip() if c.value is not None else '' for c in next(ws.iter_rows())]
idx = {h: i for i, h in enumerate(headers)}
required = ['Descripción', 'Marca', 'N° de artículo (Part)', 'GTIN']
if any(h not in idx for h in required):
    raise ValueError(f'Faltan columnas requeridas: {required}; encontradas: {headers}')

seen = set()
new_products = []
updated_products = []
invalid_rows = []
duplicate_rows = 0
for row_no, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
    if not any(v not in (None, '') for v in row):
        continue
    sku = norm_sku(row[idx['N° de artículo (Part)']])
    name = str(row[idx['Descripción']] or 'Bombilla').strip() or 'Bombilla'
    brand = str(row[idx['Marca']] or 'DT Spare Parts').strip() or 'DT Spare Parts'
    raw_gtins = re.split(r'[|;,\s]+', str(row[idx['GTIN']] or ''))
    gtins = []
    for raw in raw_gtins:
        g = norm_gtin(raw)
        if len(g) in (8, 12, 13, 14) and g not in gtins:
            gtins.append(g)
    if not sku:
        invalid_rows.append({'row': row_no, 'reason': 'missing sku'})
        continue
    if sku in seen:
        duplicate_rows += 1
        continue
    seen.add(sku)
    existing = by_sku.get(sku)
    if gtins:
        entry = gtin_map.setdefault(sku, {'sku': sku, 'gtins': []})
        entry['gtins'] = sorted(set(entry.get('gtins', []) + gtins))
    image_key = sku.replace('.', '_')
    image_source = f'0355/{image_key}_0.jpg'
    if existing:
        existing['name'] = name or existing.get('name', 'Bombilla')
        existing['description'] = f'{brand} {sku} {name}'.strip()
        existing['category'] = existing.get('category') or 'Sistema eléctrico'
        existing['subcategory'] = existing.get('subcategory') or 'Bombillas'
        existing['brand'] = brand
        existing['brandId'] = existing.get('brandId') or 'DT'
        existing['image'] = existing.get('image') or f'https://dieseltechnic.tiny.pictures/{image_source}?format=webp&width=640'
        existing['imageFull'] = existing.get('imageFull') or f'https://dieseltechnic.tiny.pictures/{image_source}?format=webp&width=1920'
        existing['imageSource'] = existing.get('imageSource') or image_source
        updated_products.append(sku)
        continue
    product = {
        'id': f'DT-{sku}', 'sku': sku, 'name': name,
        'description': f'{brand} {sku} {name}'.strip(),
        'slug': f'dt-{slugify(sku)}-{slugify(name)}',
        'url': f'https://partnerportal.dieseltechnic.com/es/products/dt-{sku.replace(".", "-")}-{slugify(name)}',
        'image': f'https://dieseltechnic.tiny.pictures/{image_source}?format=webp&width=640',
        'imageFull': f'https://dieseltechnic.tiny.pictures/{image_source}?format=webp&width=1920',
        'imageSource': image_source, 'application': 'Multibrand', 'applicationId': '0',
        'brand': brand, 'brandId': 'DT', 'manufacturer': 'Multibrand',
        'category': 'Sistema eléctrico', 'subcategory': 'Bombillas', 'usage': ['Multibrand'],
        'replaces': '', 'mainOe': '', 'packagingAmount': 1, 'salesUnit': 'PCS',
        'badges': [], 'isProductNews': False, 'isProductPromotion': False, 'crossReferences': []
    }
    new_products.append(product)
    by_sku[sku] = product

merged = list(catalog) + new_products
CATALOG_OUT.write_text(json.dumps(merged, ensure_ascii=False, separators=(',', ':')), encoding='utf-8')
GTIN_OUT.write_text(json.dumps(gtin_map, ensure_ascii=False, separators=(',', ':')), encoding='utf-8')
report = {
    'source': XLSX.name, 'sheet': ws.title, 'sourceRows': ws.max_row - 1,
    'newProducts': len(new_products), 'updatedExisting': len(updated_products),
    'duplicateRows': duplicate_rows, 'invalidRows': invalid_rows,
    'catalogBefore': len(catalog), 'catalogAfter': len(merged),
    'gtinMapBefore': len(json.loads(GTIN_IN.read_text(encoding='utf-8'))),
    'gtinMapAfter': len(gtin_map), 'newGtinsRows': sum(1 for p in new_products if gtin_map.get(p['sku'], {}).get('gtins')),
    'outputs': [str(CATALOG_OUT), str(GTIN_OUT)]
}
REPORT_OUT.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps(report, ensure_ascii=False))
