from pathlib import Path
from openpyxl import load_workbook
import json, re

xlsx = Path('/home/ubuntu/upload/Valvulas_DieselTechnic_1.xlsx')
map_in = Path('/home/ubuntu/webdev-static-assets/diesel-gtin-map-bombillas-merged.json')
map_out = Path('/home/ubuntu/webdev-static-assets/diesel-gtin-map-valvulas-xlsx_20260904.json')
catalog_in = Path('/home/ubuntu/webdev-static-assets/catalog-with-valvulas_ba8546a7.json')
catalog_out = Path('/home/ubuntu/webdev-static-assets/catalog-with-valvulas-gtin-xlsx_20260904.json')

def sku_key(v): return str(v or '').strip().upper().replace(' ', '')
def codes(v): return list(dict.fromkeys(re.findall(r'\d{8,14}', str(v or ''))))

def normalize_map(data):
    if isinstance(data, dict) and 'gtins' in data and isinstance(data['gtins'], dict): return data['gtins']
    return data

base = normalize_map(json.loads(map_in.read_text()))
merged = {sku_key(k): dict(v) for k,v in base.items()}
for key, value in merged.items():
    value['sku'] = key
    value['gtins'] = list(dict.fromkeys(str(x) for x in (value.get('gtins') or []) if x))

wb = load_workbook(xlsx, read_only=True, data_only=True)
ws = wb['Valvulas']
added_codes = 0
conflicts = []
valid_rows = 0
for desc, brand, sku, raw in ws.iter_rows(min_row=2, values_only=True):
    key = sku_key(sku); new_codes = codes(raw)
    if not key or not new_codes: continue
    valid_rows += 1
    entry = merged.setdefault(key, {'sku': key, 'gtins': []})
    old = set(entry.get('gtins') or [])
    for code in new_codes:
        if code not in old:
            entry.setdefault('gtins', []).append(code)
            added_codes += 1
    entry['gtins'] = list(dict.fromkeys(entry['gtins']))

# Detect any GTIN assigned to multiple SKUs after union; retain assignments but report for review.
occurrences = {}
for sku, entry in merged.items():
    for code in entry.get('gtins', []): occurrences.setdefault(code, []).append(sku)
conflicts = [{'gtin': code, 'skus': skus} for code, skus in occurrences.items() if len(skus) > 1]
map_out.write_text(json.dumps(merged, ensure_ascii=False, separators=(',', ':')))

catalog = json.loads(catalog_in.read_text())
valvula_skus = {sku_key(x.get('sku')) for x in catalog if str(x.get('id', '')).startswith('valvulas-')}
for item in catalog:
    key = sku_key(item.get('sku'))
    if key in valvula_skus:
        item['id'] = key
        item['internalReference'] = f"{key} / GTIN {' · '.join(merged.get(key, {}).get('gtins', []))}" if merged.get(key, {}).get('gtins') else key
    if key in merged and merged[key].get('gtins'):
        item['gtins'] = merged[key]['gtins']
catalog_out.write_text(json.dumps(catalog, ensure_ascii=False, separators=(',', ':')))
print(json.dumps({'valid_excel_rows': valid_rows, 'map_before': len(base), 'map_after': len(merged), 'new_gtin_assignments': added_codes, 'conflict_gtins': len(conflicts), 'conflict_sample': conflicts[:10], 'catalog_count': len(catalog), 'map_out': str(map_out), 'catalog_out': str(catalog_out)}, ensure_ascii=False))
