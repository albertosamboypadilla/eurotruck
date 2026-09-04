import json
from collections import Counter
from pathlib import Path
from openpyxl import load_workbook

excel = Path('/home/ubuntu/upload/Junta_Torica_DieselTechnic_1.xlsx')
catalog = Path('/home/ubuntu/webdev-static-assets/catalog-with-piston-rings-20260904.json')
gtin_map_path = Path('/home/ubuntu/webdev-static-assets/diesel-gtin-map-piston-rings-20260904.json')
rows = list(load_workbook(excel, read_only=True, data_only=True).active.iter_rows(min_row=2, values_only=True))
source = {}
for description, brand, part, gtin_cell in rows:
    sku = str(part or '').strip()
    gtins = [x.strip() for x in str(gtin_cell or '').replace(',', '|').split('|') if x.strip().isdigit()]
    if sku:
        source[sku] = {'description': str(description or '').strip(), 'brand': str(brand or '').strip(), 'gtins': sorted(set(gtins))}
items = json.loads(catalog.read_text())
by_sku = {str(item.get('sku') or '').strip(): item for item in items}
gtin_map = json.loads(gtin_map_path.read_text())
missing = []
new_gtins = []
conflicts = []
for sku, row in source.items():
    item = by_sku.get(sku)
    existing = set((item or {}).get('gtins') or []) | set((gtin_map.get(sku) or {}).get('gtins') or [])
    if not item:
        missing.append(sku)
    if set(row['gtins']) - existing:
        new_gtins.append({'sku': sku, 'new': sorted(set(row['gtins']) - existing), 'existing': sorted(existing)})
    for gtin in row['gtins']:
        owners = [candidate for candidate in items if gtin in (candidate.get('gtins') or []) and str(candidate.get('sku') or '') != sku]
        if owners:
            conflicts.append({'sku': sku, 'gtin': gtin, 'owners': [str(owner.get('sku') or '') for owner in owners[:5]]})
print(json.dumps({'source_rows': len(rows), 'unique_skus': len(source), 'catalog_count': len(items), 'missing_skus': missing, 'new_gtin_rows': new_gtins[:25], 'new_gtin_row_count': len(new_gtins), 'conflicts': conflicts[:25], 'conflict_count': len(conflicts)}, ensure_ascii=False, indent=2))
