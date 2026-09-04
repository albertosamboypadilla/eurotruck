from openpyxl import load_workbook
from pathlib import Path
import json

path = Path('/home/ubuntu/upload/EUROTRUCK_Reporte_Inventario_2026-09-03.xlsx')
ws = load_workbook(path, data_only=False).active
nonempty = []
for row in ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True):
    values = list(row)
    if any(value not in (None, '') for value in values):
        nonempty.append(values)
summary = {
    'nonempty_rows': len(nonempty),
    'first_20_nonempty_rows': nonempty[:20],
    'last_10_nonempty_rows': nonempty[-10:],
    'header_rows': [list(ws.iter_rows(min_row=row, max_row=row, values_only=True))[0] for row in range(1, min(ws.max_row, 8) + 1)],
    'body_row_styles': {
        'row4': {'fill': ws['A4'].fill.fgColor.rgb, 'font_size': ws['A4'].font.sz, 'bold': ws['A4'].font.bold, 'border_left': ws['A4'].border.left.style, 'border_bottom': ws['A4'].border.bottom.style},
        'row5': {'fill': ws['A5'].fill.fgColor.rgb, 'font_size': ws['A5'].font.sz, 'bold': ws['A5'].font.bold, 'border_left': ws['A5'].border.left.style, 'border_bottom': ws['A5'].border.bottom.style},
    },
}
print(json.dumps(summary, ensure_ascii=False, indent=2, default=str))
