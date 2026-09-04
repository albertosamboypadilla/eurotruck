from pathlib import Path
from openpyxl import load_workbook
import json, re

xlsx = Path('/home/ubuntu/upload/Valvulas_DieselTechnic_1.xlsx')
catalog_path = Path('/home/ubuntu/webdev-static-assets/catalog-with-valvulas_ba8546a7.json')
map_candidates = [Path('/home/ubuntu/webdev-static-assets/diesel-gtin-map-bombillas-merged.json'), Path('/home/ubuntu/webdev-static-assets/diesel-gtin-map-complete.json'), Path('/home/ubuntu/webdev-static-assets/diesel-gtin-map-bombillas-merged_7d3b08a9.json'), Path('/tmp/diesel-gtin-map-bombillas-merged_7d3b08a9.json')]
map_path = next((p for p in map_candidates if p.exists()), None)
wb = load_workbook(xlsx, read_only=True, data_only=True)
ws = wb['Valvulas']
rows = list(ws.iter_rows(min_row=2, values_only=True))

def normalize_sku(v):
    return str(v or '').strip().upper().replace(' ', '')

def gtins(v):
    vals = re.findall(r'\d{8,14}', str(v or ''))
    return list(dict.fromkeys(vals))

catalog = json.loads(catalog_path.read_text())
catalog_skus = {normalize_sku(x.get('sku')) for x in catalog}
map_data = json.loads(map_path.read_text()) if map_path else {}
if isinstance(map_data, dict) and 'gtins' in map_data: map_data = map_data['gtins']
map_skus = {normalize_sku(k) for k in map_data} if isinstance(map_data, dict) else set()
valid=[]; invalid=[]; duplicate_skus={}; duplicate_gtins={}; seen_gtins={}
for row in rows:
    desc, brand, sku, raw = row[:4]
    key=normalize_sku(sku); codes=gtins(raw)
    if not key or not codes:
        invalid.append({'sku':sku,'gtin':raw,'reason':'missing_sku_or_gtin'}); continue
    duplicate_skus.setdefault(key,0); duplicate_skus[key]+=1
    for code in codes: duplicate_gtins.setdefault(code,0); duplicate_gtins[code]+=1; seen_gtins.setdefault(code,[]).append(key)
    valid.append({'sku':key,'description':str(desc or '').strip(),'brand':str(brand or '').strip(),'gtins':codes})
report={'excel_rows':len(rows),'valid_rows':len(valid),'invalid_rows':len(invalid),'unique_excel_skus':len({x['sku'] for x in valid}),'catalog_matches':sum(x['sku'] in catalog_skus for x in valid),'catalog_missing':sum(x['sku'] not in catalog_skus for x in valid),'already_in_map':sum(x['sku'] in map_skus for x in valid),'duplicate_sku_rows':sum(v>1 for v in duplicate_skus.values()),'duplicate_gtin_values':sum(v>1 for v in duplicate_gtins.values()),'sample_missing':[x for x in valid if x['sku'] not in catalog_skus][:10],'sample_duplicates':[(k,v) for k,v in duplicate_skus.items() if v>1][:10], 'map_path':str(map_path) if map_path else None}
print(json.dumps(report, ensure_ascii=False, indent=2))
Path('/tmp/valvulas-gtin-analysis.json').write_text(json.dumps({'report':report,'valid':valid,'invalid':invalid,'gtin_occurrences':seen_gtins},ensure_ascii=False,indent=2))
