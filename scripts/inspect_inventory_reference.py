from openpyxl import load_workbook
from pathlib import Path
import json

path = Path('/home/ubuntu/upload/EUROTRUCK_Reporte_Inventario_2026-09-03.xlsx')
wb = load_workbook(path, data_only=False)
ws = wb[wb.sheetnames[0]]

def color(value):
    if value is None:
        return None
    return {'type': value.type, 'rgb': value.rgb, 'indexed': value.indexed, 'theme': value.theme, 'tint': value.tint}

summary = {
    'sheetnames': wb.sheetnames,
    'active_sheet': ws.title,
    'dimensions': ws.calculate_dimension(),
    'max_row': ws.max_row,
    'max_column': ws.max_column,
    'freeze_panes': str(ws.freeze_panes),
    'merged_cells': [str(value) for value in ws.merged_cells.ranges],
    'headers': [ws.cell(1, col).value for col in range(1, ws.max_column + 1)],
    'sample_rows': [[ws.cell(row, col).value for col in range(1, ws.max_column + 1)] for row in range(2, min(ws.max_row, 6) + 1)],
    'column_widths': {letter: dimension.width for letter, dimension in ws.column_dimensions.items() if dimension.width is not None},
    'row_heights': {str(row): dimension.height for row, dimension in ws.row_dimensions.items() if dimension.height is not None},
    'autofilter': str(ws.auto_filter.ref),
    'print_title_rows': str(ws.print_title_rows),
    'print_area': str(ws.print_area),
    'orientation': ws.page_setup.orientation,
    'paper_size': ws.page_setup.paperSize,
    'fit_to_width': ws.page_setup.fitToWidth,
    'fit_to_height': ws.page_setup.fitToHeight,
    'first_row_style': {
        'fill': color(ws.cell(1, 1).fill.fgColor),
        'font': {'name': ws.cell(1, 1).font.name, 'size': ws.cell(1, 1).font.sz, 'bold': ws.cell(1, 1).font.bold, 'color': color(ws.cell(1, 1).font.color)},
        'alignment': {'horizontal': ws.cell(1, 1).alignment.horizontal, 'vertical': ws.cell(1, 1).alignment.vertical, 'wrap_text': ws.cell(1, 1).alignment.wrap_text},
    },
    'body_style': {
        'fill': color(ws.cell(2, 1).fill.fgColor),
        'font': {'name': ws.cell(2, 1).font.name, 'size': ws.cell(2, 1).font.sz, 'bold': ws.cell(2, 1).font.bold, 'color': color(ws.cell(2, 1).font.color)},
        'alignment': {'horizontal': ws.cell(2, 1).alignment.horizontal, 'vertical': ws.cell(2, 1).alignment.vertical, 'wrap_text': ws.cell(2, 1).alignment.wrap_text},
    },
}
Path('/tmp/inventory-reference-summary.json').write_text(json.dumps(summary, ensure_ascii=False, indent=2, default=str))
print(json.dumps(summary, ensure_ascii=False, indent=2, default=str))
