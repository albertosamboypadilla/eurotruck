from pathlib import Path
from openpyxl import load_workbook
import json

path = Path('/home/ubuntu/upload/Valvulas_DieselTechnic_1.xlsx')
wb = load_workbook(path, read_only=True, data_only=True)
report = {'sheets': wb.sheetnames, 'rows': []}
for ws in wb.worksheets:
    rows = ws.iter_rows(values_only=True)
    header = next(rows, ())
    header = [str(v).strip() if v is not None else '' for v in header]
    report['rows'].append({'sheet': ws.title, 'max_row': ws.max_row, 'max_column': ws.max_column, 'header': header, 'sample': [[v for v in row] for _, row in zip(range(5), rows)]})
print(json.dumps(report, ensure_ascii=False, default=str, indent=2))
